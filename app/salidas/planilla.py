"""Planillas: Excel (.xlsx) y CSV con los registros, para trabajarlos en una computadora.

Decisión de producto "Salidas": "Excel" o "planilla" -> .xlsx; CSV solo si lo pide, con `;`
y BOM (así Excel en castellano lo abre en columnas y con las tildes bien).

El texto que escribió la persona nunca se toma como fórmula: en Excel cada texto se guarda
como texto, y en el CSV un texto que empieza con =, +, - o @ lleva un apóstrofo adelante
(si no, Excel lo ejecutaría al abrirlo).
"""

import csv
import io
from collections.abc import Sequence
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.infra.marca import NOMBRE
from app.salidas.pdf import PIE, Membrete

Valor = str | int | float | date | None

_FORMULA = ("=", "+", "-", "@", "\t", "\r")
VERDE = "2E7D32"


def _texto_seguro(texto: str) -> str:
    return "'" + texto if texto.startswith(_FORMULA) else texto


def _numero_csv(valor: float) -> str:
    """Con coma decimal, como lo espera Excel en castellano: 2.5 -> "2,5"; 20.0 -> "20"."""
    return f"{valor:g}".replace(".", ",") if isinstance(valor, float) else str(valor)


def csv_bytes(columnas: Sequence[str], filas: Sequence[Sequence[Valor]]) -> bytes:
    salida = io.StringIO()
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(columnas)
    for fila in filas:
        celdas = []
        for valor in fila:
            if valor is None:
                celdas.append("")
            elif isinstance(valor, date):
                celdas.append(valor.strftime("%d/%m/%Y"))
            elif isinstance(valor, int | float):
                celdas.append(_numero_csv(valor))
            else:
                celdas.append(_texto_seguro(valor))
        escritor.writerow(celdas)
    return salida.getvalue().encode("utf-8-sig")  # con BOM


def excel_bytes(
    membrete: Membrete,
    titulo: str,
    subtitulo: list[str],
    resumen: list[str],
    columnas: Sequence[str],
    filas: Sequence[Sequence[Valor]],
    generado: datetime,
) -> bytes:
    """Una hoja "Registros" con la tabla (lista para filtrar y sumar) y una "Resumen" con el
    membrete, de qué es, los totales y el pie legal."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Registros"
    hoja.append(list(columnas))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=VERDE)
    for fila in filas:
        hoja.append(list(fila))
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                celda.data_type = "s"  # texto, nunca fórmula
            elif isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for i, columna in enumerate(columnas, start=1):
        largo = max([len(columna), *(len(str(f[i - 1])) for f in filas if f[i - 1] is not None)])
        hoja.column_dimensions[get_column_letter(i)].width = min(max(largo + 2, 10), 45)

    resumen_hoja = libro.create_sheet("Resumen")
    lineas: list[str] = [
        NOMBRE,
        titulo,
        membrete.linea(),
        " · ".join(subtitulo),
        f"Generado el {generado:%d/%m/%Y %H:%M}",
    ]
    if resumen:
        lineas += ["", "Totales", *resumen]
    lineas += ["", PIE]
    for linea in lineas:
        resumen_hoja.append([linea])
        resumen_hoja.cell(resumen_hoja.max_row, 1).data_type = "s"
    resumen_hoja["A1"].font = Font(bold=True, size=14, color=VERDE)
    resumen_hoja["A2"].font = Font(bold=True, size=12)
    resumen_hoja.column_dimensions["A"].width = 100

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
