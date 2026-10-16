"""Archivos: los registros en PDF, Excel o CSV, enviados por WhatsApp como documento."""

import io
import json
from datetime import UTC, date, datetime

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from app import main
from app.api import webhook
from app.canales import whatsapp
from app.ia import interprete
from app.infra.config import config
from app.modulos.comun import exportar
from app.nucleo import pipeline
from app.salidas import pdf
from tests.ayudas import con_perfil, con_perfil_en

HOY = date(2026, 10, 15)
TEL = "5491100000001"


def _evento(tipo: str, **datos) -> dict:
    categoria = "campo" if tipo == "lluvia" else "sanidad"
    return {"intencion": "registrar", "eventos": [{"categoria": categoria, "tipo": tipo, **datos}]}


@pytest.fixture
def ia(monkeypatch):
    respuestas: dict[str, object] = {
        "estoy en la pastora": {"intencion": "establecimiento", "establecimiento": "La Pastora"},
        "vacuné 30 vaquillonas": _evento(
            "vacunacion", producto="aftosa", cantidad_animales=30, categoria_animal="vaquillonas", lote="3"
        ),
        "llovieron 20 mm": _evento("lluvia", valor=20, unidad="mm", fecha_evento="2026-10-14"),
        "en el ombú desparasité 10 terneros": {
            **_evento("desparasitacion", producto="ivermectina", cantidad_animales=10, categoria_animal="terneros"),
            "establecimiento": "El Ombú",
        },
        "mandame un pdf": {"intencion": "exportar", "exportar": {"formato": "pdf"}, "consulta": {}},
        "pdf de la pastora": {"intencion": "exportar", "consulta": {"establecimiento": "la pastora"}},
    }
    monkeypatch.setattr(pipeline, "_ahora", lambda: datetime(2026, 10, 15, 15, 30, tzinfo=UTC))
    monkeypatch.setattr(interprete, "interpretar", lambda texto, *_a, **_k: respuestas[texto])
    return respuestas


def _cargar(db) -> None:
    usuario = con_perfil(db, TEL, rubro="veterinario")
    usuario.matricula = "MP 1234"
    for texto in (
        "estoy en la pastora",
        "vacuné 30 vaquillonas",
        "llovieron 20 mm",
        "en el ombú desparasité 10 terneros",
    ):
        pipeline.manejar_mensaje(db, TEL, texto, hoy=HOY)


def _texto(contenido: bytes) -> str:
    paginas = PdfReader(io.BytesIO(contenido)).pages
    return " ".join(" ".join(p.extract_text() for p in paginas).split())


def test_un_pdf_con_todo_lo_anotado(db, ia):
    _cargar(db)
    r = pipeline.atender(db, TEL, "mandame un pdf", hoy=HOY)

    assert r.texto.startswith("📄 *PDF LISTO*\nTodo lo anotado · 3 registros\n")
    [adjunto] = r.adjuntos
    assert (adjunto.nombre, adjunto.tipo) == ("CampoBot-registros-2026-10-15.pdf", "application/pdf")
    assert adjunto.contenido.startswith(b"%PDF")

    texto = _texto(adjunto.contenido)
    assert "CampoBot" in texto and "Generado el 15/10/2026 12:30" in texto  # hora de Argentina
    assert "Registro de trabajo Juana Gómez · Veterinario · Tandil · Matrícula MP 1234" in texto
    assert "Totales" in texto and "Vacunación: 30 animales" in texto and "Lluvia: 20 mm" in texto
    assert texto.index("14/10/2026") < texto.index("15/10/2026 La Pastora")  # del más viejo al más nuevo
    assert "Vacunación · aftosa 30 vaquillonas · Lote 3" in texto
    assert "El Ombú Desparasitación · ivermectina 10 terneros" in texto
    assert " ".join(pdf.PIE.split()) in texto
    assert "Página 1 de 1" in texto


def test_un_pdf_de_un_campo(db, ia):
    _cargar(db)
    r = pipeline.atender(db, TEL, "pdf de la pastora", hoy=HOY)
    assert r.texto.startswith("📄 *PDF LISTO*\nCampo: La Pastora · 2 registros\n")
    texto = _texto(r.adjuntos[0].contenido)
    assert "La Pastora" in texto and "Ombú" not in texto


def test_sin_registros_no_manda_nada(db, ia):
    con_perfil(db, TEL)
    r = pipeline.atender(db, TEL, "mandame un pdf", hoy=HOY)
    assert (r.texto, r.adjuntos) == (exportar.SIN_RESULTADOS, [])


def test_el_texto_de_la_persona_no_se_interpreta():
    """Ni markdown ni HTML: va tal cual. Lo que Helvetica no tiene (emojis) se saca."""
    tabla = pdf.Tabla(("A", "B"), (1, 1), [["**negrita** <b>x</b> 🐄", "Lote 3 → 5"]])
    contenido = pdf.documento(pdf.Membrete("Ana Paz", "Productor", "Azul"), "Título", [], [], tabla, datetime.now())
    texto = _texto(contenido)
    assert "**negrita** <b>x</b>" in texto and "Lote 3 -> 5" in texto
    assert pdf.latin1("Ñandú 🐄 “sí” – ok") == 'Ñandú "sí" - ok'


def test_muchas_filas_ocupan_varias_paginas():
    tabla = pdf.Tabla(("Fecha", "Detalle"), (1, 3), [[f"{i:02d}/10/2026", "x" * 40] for i in range(150)])
    contenido = pdf.documento(pdf.Membrete("Ana Paz", "Productor", "Azul"), "Título", [], [], tabla, datetime.now())
    paginas = PdfReader(io.BytesIO(contenido)).pages
    assert len(paginas) > 2
    assert all(" ".join(pdf.PIE.split())[:40] in " ".join(p.extract_text().split()) for p in paginas)


# --- Excel y CSV ----------------------------------------------------------------------


def _con_formula(db, ia) -> None:
    """Un producto que parece una fórmula: no se tiene que ejecutar al abrir la planilla."""
    _cargar(db)
    ia["traté 2 toros con =1+1"] = _evento("tratamiento", producto="=1+1", cantidad_animales=2, valor=2.5, unidad="ml")
    pipeline.manejar_mensaje(db, TEL, "traté 2 toros con =1+1", hoy=HOY)


def test_una_planilla_excel(db, ia):
    _con_formula(db, ia)
    ia["pasame la planilla"] = {"intencion": "exportar", "exportar": {"formato": "excel"}, "consulta": {}}
    r = pipeline.atender(db, TEL, "pasame la planilla", hoy=HOY)

    assert r.texto.startswith("📊 *PLANILLA LISTA*\nTodo lo anotado · 4 registros\n")
    [adjunto] = r.adjuntos
    assert adjunto.nombre == "CampoBot-registros-2026-10-15.xlsx" and adjunto.tipo == exportar.XLSX
    libro = load_workbook(io.BytesIO(adjunto.contenido))
    assert libro.sheetnames == ["Registros", "Resumen"]
    filas = list(libro["Registros"].iter_rows(values_only=True))
    assert filas[0] == exportar.COLUMNAS_PLANILLA
    lluvia = filas[1]  # la más vieja primero
    assert (lluvia[0], lluvia[1], lluvia[5], lluvia[11], lluvia[12]) == (
        datetime(2026, 10, 14),
        "La Pastora",
        "Lluvia",
        20,
        "mm",
    )
    vacunacion = next(f for f in filas if f[5] == "Vacunación")
    assert (vacunacion[2], vacunacion[6], vacunacion[8], vacunacion[9]) == ("3", "aftosa", 30, "vaquillonas")
    formula = next(f for f in filas if f[5] == "Tratamiento")
    assert formula[6] == "=1+1" and formula[11] == 2.5
    celda = next(c for c in libro["Registros"]["G"] if c.value == "=1+1")
    assert celda.data_type == "s"  # texto, no fórmula

    resumen = [f[0] for f in libro["Resumen"].iter_rows(values_only=True)]
    assert resumen[:3] == ["CampoBot", "Registro de trabajo", "Juana Gómez · Veterinario · Tandil · Matrícula MP 1234"]
    assert "Vacunación: 30 animales" in resumen and resumen[-1] == pdf.PIE


def test_un_csv_con_punto_y_coma_y_bom(db, ia):
    _con_formula(db, ia)
    ia["mandame un csv"] = {"intencion": "exportar", "exportar": {"formato": "csv"}, "consulta": {}}
    r = pipeline.atender(db, TEL, "mandame un csv", hoy=HOY)

    assert r.texto.startswith("📊 *CSV LISTO*\n")
    [adjunto] = r.adjuntos
    assert adjunto.nombre == "CampoBot-registros-2026-10-15.csv" and adjunto.tipo == "text/plain"
    assert adjunto.contenido.startswith(b"\xef\xbb\xbf")  # BOM
    lineas = adjunto.contenido.decode("utf-8-sig").splitlines()
    assert lineas[0] == ";".join(exportar.COLUMNAS_PLANILLA)
    assert lineas[1] == "14/10/2026;La Pastora;;;campo;Lluvia;;;;;;20;mm;"
    assert "'=1+1" in lineas[-1] and ";2,5;ml;" in lineas[-1]


def test_sin_formato_es_pdf(db, ia):
    _cargar(db)
    ia["mandame lo anotado"] = {"intencion": "exportar", "exportar": {"formato": "word"}}
    assert pipeline.atender(db, TEL, "mandame lo anotado", hoy=HOY).adjuntos[0].nombre.endswith(".pdf")


# --- Envío por WhatsApp ---------------------------------------------------------------


def test_el_pdf_llega_despues_del_texto(servidor, motor, ia, monkeypatch):
    con_perfil_en(motor, TEL)
    documentos: list[tuple[str, str, str]] = []

    def enviar_documento(telefono, nombre, contenido, tipo):
        documentos.append((telefono, nombre, tipo))
        servidor.append((telefono, f"[{nombre}]"))
        return True

    monkeypatch.setattr(whatsapp, "enviar_documento", enviar_documento)
    monkeypatch.setattr(pipeline, "_hoy", lambda: HOY)
    with TestClient(main.app) as c:
        for i, texto in enumerate(("vacuné 30 vaquillonas", "mandame un pdf")):
            mensaje = {"from": TEL, "id": f"wamid.{i}", "type": "text", "text": {"body": texto}, "timestamp": str(i)}
            c.post("/webhook", json={"entry": [{"changes": [{"value": {"messages": [mensaje]}}]}]})
            assert webhook.fila.esperar(5)

    assert documentos == [(TEL, "CampoBot-registros-2026-10-15.pdf", "application/pdf")]
    assert servidor[-2][1].startswith("📄 *PDF LISTO*") and servidor[-1][1] == "[CampoBot-registros-2026-10-15.pdf]"


def test_si_el_archivo_no_sale_avisa(servidor, motor, ia, monkeypatch):
    con_perfil_en(motor, TEL)
    monkeypatch.setattr(whatsapp, "enviar_documento", lambda *_a: False)
    with TestClient(main.app) as c:
        for i, texto in enumerate(("vacuné 30 vaquillonas", "mandame un pdf")):
            mensaje = {"from": TEL, "id": f"wamid.{i}", "type": "text", "text": {"body": texto}, "timestamp": str(i)}
            c.post("/webhook", json={"entry": [{"changes": [{"value": {"messages": [mensaje]}}]}]})
            assert webhook.fila.esperar(5)
    assert servidor[-1] == (TEL, webhook.ERROR_ARCHIVO)


class _Resp:
    def __init__(self, status: int, cuerpo: dict) -> None:
        self.status_code, self._cuerpo, self.text = status, cuerpo, json.dumps(cuerpo)

    def json(self) -> dict:
        return self._cuerpo


def test_enviar_documento_sube_el_archivo_y_lo_manda(monkeypatch):
    llamadas: list[tuple[str, dict]] = []

    def post(url, headers, timeout, json=None, data=None, files=None):
        llamadas.append((url.rsplit("/", 1)[-1], json or {**(data or {}), "archivo": files["file"][0]}))
        return _Resp(200, {"id": "media-1"} if url.endswith("/media") else {})

    monkeypatch.setattr(config, "wa_ar_quitar_9", False)
    monkeypatch.setattr(httpx, "post", post)
    assert whatsapp.enviar_documento(TEL, "informe.pdf", b"%PDF-1.3", "application/pdf")
    assert llamadas == [
        ("media", {"messaging_product": "whatsapp", "type": "application/pdf", "archivo": "informe.pdf"}),
        (
            "messages",
            {
                "messaging_product": "whatsapp",
                "to": TEL,
                "type": "document",
                "document": {"id": "media-1", "filename": "informe.pdf"},
            },
        ),
    ]


def test_si_no_se_puede_subir_no_lo_manda(monkeypatch):
    monkeypatch.setattr(httpx, "post", lambda url, **_k: _Resp(400, {"error": {"code": 100}}))
    assert not whatsapp.enviar_documento(TEL, "informe.pdf", b"%PDF", "application/pdf")
